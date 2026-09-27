"""Verify exported review data against CSV sources and preserve open gates."""
from pathlib import Path
import csv,json,re,math,zipfile
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[3]
P=ROOT/'outputs/research/review_package_v1.8'
wb=load_workbook(P/'supplementary_data_1_v1.8_review.xlsx',read_only=True,data_only=True)
tables={'Evidence gates':'evidence_gates.csv','Conditional factorials':'conditional_factorials.csv','External power':'external_power_corrected.csv','Provincial inputs':'provincial_input_register.csv'}
counts={};checked=0
for sheet,file in tables.items():
 rows=list(csv.reader((P/'data'/file).open()))
 for i,row in enumerate(rows,1):
  for j,expected in enumerate(row,1):
   actual=wb[sheet].cell(i,j).value
   try:
    number=float(expected)
   except ValueError:
    assert actual==expected,(sheet,i,j,actual,expected)
   else:
    assert isinstance(actual,(int,float)) and math.isclose(actual,number,rel_tol=1e-13,abs_tol=1e-12),(sheet,i,j)
   checked+=1
 counts[sheet]=len(rows)-1
assert counts=={'Evidence gates':6,'Conditional factorials':108,'External power':24,'Provincial inputs':10}
for row in csv.DictReader((P/'data/provincial_input_register.csv').open()):
 assert (ROOT/'outputs/research'/row['evidence_path_from_outputs_research']).is_file()
for f in (P/'field_intake').glob('*.csv'):
 assert len(list(csv.reader(f.open())))==1,'Field intake contains unexpected observation'
gates=list(csv.DictReader((P/'data/evidence_gates.csv').open()))
assert all(r['status'] not in ['complete','pass'] for r in gates)
md=(P/'core_paper_en_v1.8_review.md').read_text()
body,tail=md.split('## References\n',1)
refs=tail.split('## Evidence register')[0]
numbers=[int(x) for x in re.findall(r'(?m)^(\d+)\. ',refs)]
assert numbers==list(range(1,27))
cited=[]
for match in re.findall(r'\[([\d,–-]+)\]',body):
 for part in match.split(','):
  if re.search('[–-]',part):
   a,b=map(int,re.split('[–-]',part));cited.extend(range(a,b+1))
  else:cited.append(int(part))
assert list(dict.fromkeys(cited))==numbers
assert '4.115–5.960 kW' in body and '3.777–5.377 kW' not in body
assert 'not submission-ready' in body
for mdpath in P.glob('*.md'):
 for image in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',mdpath.read_text()):
  assert (P/image).exists()
pages={}
for f in P.glob('*.docx'):
 with zipfile.ZipFile(f) as z:
  xml=z.read('word/document.xml').decode()
  assert '<w:tbl>' not in xml or '<w:cantSplit' in xml
 pdf=P/(f.stem+'.pdf')
 assert pdf.exists()
 pages[pdf.name]=len(PdfReader(pdf).pages)
out=dict(status='pass',cells_compared_with_csv=checked,table_rows=counts,
    references=26,pdf_pages=pages,field_observations=0,provincial_holdout_executed=False,
    submission_ready=False,scope='Export integrity and evidence-boundary verification, not empirical acceptance')
(P/'package_validation.json').write_text(json.dumps(out,indent=2)+'\n')
print(json.dumps(out,indent=2))
