#!/usr/bin/env python3
"""Build the current review manuscript and supplement as editable DOCX files."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
PACKAGE = ROOT / "outputs/research/current_paper_20260927/01_current_revision"

EQUATIONS = {
    r"M = \frac{(C_{00}-C_{10})+(C_{01}-C_{11})}{2}": "mode",
    r"T = \frac{(C_{00}-C_{01})+(C_{10}-C_{11})}{2}": "start",
    r"\Delta E = E_r-E_m-K+P_0(T_m+\tau-T_r)": "energy",
    r"D^p=B+A^p": "demand",
    r"B=L-A^r": "background",
}


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def m_run(text, italic=True):
    run = OxmlElement("m:r")
    props = OxmlElement("m:rPr")
    if not italic:
        sty = OxmlElement("m:sty")
        sty.set(qn("m:val"), "p")
        props.append(sty)
    run.append(props)
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def m_sub(base, subscript):
    node = OxmlElement("m:sSub")
    node.append(OxmlElement("m:sSubPr"))
    e = OxmlElement("m:e")
    e.append(m_run(base))
    sub = OxmlElement("m:sub")
    sub.append(m_run(subscript, italic=False))
    node.extend((e, sub))
    return node


def m_sup(base, superscript):
    node = OxmlElement("m:sSup")
    node.append(OxmlElement("m:sSupPr"))
    e = OxmlElement("m:e")
    e.append(m_run(base))
    sup = OxmlElement("m:sup")
    sup.append(m_run(superscript, italic=False))
    node.extend((e, sup))
    return node


def m_frac(numerator_nodes, denominator_nodes):
    node = OxmlElement("m:f")
    node.append(OxmlElement("m:fPr"))
    num = OxmlElement("m:num")
    for item in numerator_nodes:
        num.append(item)
    den = OxmlElement("m:den")
    for item in denominator_nodes:
        den.append(item)
    node.extend((num, den))
    return node


def m_paren(nodes):
    node = OxmlElement("m:d")
    props = OxmlElement("m:dPr")
    beg = OxmlElement("m:begChr")
    beg.set(qn("m:val"), "(")
    end = OxmlElement("m:endChr")
    end.set(qn("m:val"), ")")
    props.extend((beg, end))
    element = OxmlElement("m:e")
    for item in nodes:
        element.append(item)
    node.extend((props, element))
    return node


def cost_difference(left, right):
    return m_paren([m_sub("C", left), m_run("−", italic=False), m_sub("C", right)])


def add_equation(document, kind):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    equation = OxmlElement("m:oMath")
    if kind in {"mode", "start"}:
        lhs = "M" if kind == "mode" else "T"
        pairs = (("00", "10"), ("01", "11")) if kind == "mode" else (("00", "01"), ("10", "11"))
        equation.append(m_run(lhs))
        equation.append(m_run(" = ", italic=False))
        numerator = [cost_difference(*pairs[0]), m_run(" + ", italic=False), cost_difference(*pairs[1])]
        equation.append(m_frac(numerator, [m_run("2", italic=False)]))
    elif kind == "energy":
        equation.append(m_run("Δ"))
        equation.append(m_run("E"))
        equation.append(m_run(" = ", italic=False))
        equation.append(m_sub("E", "r"))
        equation.append(m_run(" − ", italic=False))
        equation.append(m_sub("E", "m"))
        equation.append(m_run(" − ", italic=False))
        equation.append(m_run("K"))
        equation.append(m_run(" + ", italic=False))
        equation.append(m_sub("P", "0"))
        equation.append(m_paren([
            m_sub("T", "m"), m_run(" + ", italic=False), m_run("τ"),
            m_run(" − ", italic=False), m_sub("T", "r")
        ]))
    elif kind == "demand":
        equation.append(m_sup("D", "p"))
        equation.append(m_run(" = ", italic=False))
        equation.append(m_run("B"))
        equation.append(m_run(" + ", italic=False))
        equation.append(m_sup("A", "p"))
    elif kind == "background":
        equation.append(m_run("B"))
        equation.append(m_run(" = ", italic=False))
        equation.append(m_run("L"))
        equation.append(m_run(" − ", italic=False))
        equation.append(m_sup("A", "r"))
    else:
        raise ValueError(kind)
    paragraph._p.append(equation)


def add_inline(paragraph, text, *, base_bold=False):
    """Add basic Markdown bold and URLs without leaving Markdown marks."""
    pattern = re.compile(r"(\*\*.+?\*\*|https?://\S+)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            run.bold = base_bold
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            cleaned = token.rstrip(".,;)")
            trailer = token[len(cleaned):]
            run = paragraph.add_run(cleaned)
            run.font.color.rgb = RGBColor(31, 78, 121)
            run.underline = True
            if trailer:
                paragraph.add_run(trailer)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        run.bold = base_bold


def parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        rows.append(cells)
        idx += 1
    if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        raise ValueError(f"Malformed Markdown table near line {start + 1}")
    return [rows[0]] + rows[2:], idx


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {
        3: [Inches(2.85), Inches(1.05), Inches(2.55)],
        4: [Inches(1.55), Inches(0.85), Inches(2.0), Inches(2.05)],
    }.get(len(rows[0]), [Inches(6.45 / len(rows[0]))] * len(rows[0]))
    for ridx, row_values in enumerate(rows):
        row = table.rows[ridx]
        if ridx == 0:
            set_repeat_table_header(row)
        for cidx, value in enumerate(row_values):
            cell = row.cells[cidx]
            cell.width = widths[cidx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if ridx == 0:
                set_cell_fill(cell, "1F4E79")
            elif ridx % 2 == 0:
                set_cell_fill(cell, "F3F6F9")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                if ridx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def configure_document(document, short_title):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11.3)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 14.5, 14, 6),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)


def build(markdown_path, output_path, short_title):
    document = Document()
    configure_document(document, short_title)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    idx = 0
    first_heading = True
    front_matter = True
    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped:
            idx += 1
            continue

        if stripped.startswith("$$"):
            chunks = []
            idx += 1
            while idx < len(lines) and not lines[idx].strip().startswith("$$"):
                chunks.append(lines[idx].strip())
                idx += 1
            if idx >= len(lines):
                raise ValueError(f"Unclosed equation in {markdown_path}")
            source = "".join(chunks)
            if source not in EQUATIONS:
                raise ValueError(f"Unknown equation: {source}")
            add_equation(document, EQUATIONS[source])
            idx += 1
            continue

        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(document, rows)
            continue

        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = markdown_path.parent / image_match.group(2)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.62))
            idx += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text_value = heading_match.group(2)
            if first_heading and level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, text_value)
                first_heading = False
            else:
                if level == 2 and text_value.startswith("Figure"):
                    document.add_page_break()
                if level == 2:
                    front_matter = False
                paragraph = document.add_paragraph(style=f"Heading {min(level - 1, 3)}")
                add_inline(paragraph, text_value)
            idx += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            idx += 1
            continue

        paragraph = document.add_paragraph()
        if front_matter and not stripped.startswith("**"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.font.name = "Arial"
        if stripped.startswith("**Manuscript status.**") or stripped.startswith("**Scope.**"):
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(10)
        if stripped.startswith("**Figure"):
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
        add_inline(paragraph, stripped)
        idx += 1

    core = document.core_properties
    core.title = lines[0].lstrip("# ")
    core.subject = "Revision v1.5 review edition; evidence through revision stage 35"
    core.author = "William Wei; Lanlan Liu"
    core.comments = "Prepared as an evidence-bounded review manuscript on 27 September 2026."
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main():
    pairs = [
        (PACKAGE / "core_paper_en_v1.5_review.md", PACKAGE / "core_paper_en_v1.5_review.docx", "AI computing flexibility · v1.5 review"),
        (PACKAGE / "supplementary_information_v1.5_review.md", PACKAGE / "supplementary_information_v1.5_review.docx", "Supplementary methods · v1.5 review"),
    ]
    for source, target, short_title in pairs:
        build(source, target, short_title)
        print(target)


if __name__ == "__main__":
    sys.exit(main())
