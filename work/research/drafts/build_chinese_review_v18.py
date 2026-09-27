#!/usr/bin/env python3
"""Build the Chinese reading edition without altering scientific outputs."""
from pathlib import Path
import shutil
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_current_paper_docx import build

ROOT = Path(__file__).resolve().parents[3]
EN = ROOT / 'outputs/research/review_package_v1.8'
ZH = ROOT / 'outputs/research/review_package_v1.8_zh'

def main():
    ZH.mkdir(exist_ok=True)
    shutil.copytree(EN / 'figures', ZH / 'figures', dirs_exist_ok=True)
    main_md = ZH / '论文主文_v1.8_中文.md'
    source = (EN / 'core_paper_en_v1.8_review.md').read_text()
    references = source.split('## References\n', 1)[1].split('## Evidence register', 1)[0].strip()
    text = main_md.read_text()
    if '<!-- REFERENCES_FROM_ENGLISH -->' in text:
        main_md.write_text(text.replace('<!-- REFERENCES_FROM_ENGLISH -->', references))
    else:
        assert text.split('## 参考文献\n',1)[1].split('## 证据登记',1)[0].strip() == references
    for stem, label in [('论文主文_v1.8_中文','论文主文'),('补充信息_v1.8_中文','补充信息'),('投稿信_v1.8_中文草稿','投稿信草稿')]:
        path = ZH / (stem + '.docx')
        build(ZH / (stem + '.md'), path, label+'  v1.8 中文审读版', subject='Chinese translation of v1.8 review edition, evidence through stage 38', comments='Not a submission-ready empirical study. English reference and figure labels retained.')
        doc = Document(path)
        for sec in doc.sections:
            sec.page_width = Mm(210); sec.page_height = Mm(297)
            sec.left_margin = sec.right_margin = Mm(19)
            sec.top_margin = Mm(20); sec.bottom_margin = Mm(19)
        for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
            st = doc.styles[name]
            st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC' if name in ('Normal','List Bullet') else 'Heiti SC')
            st.font.color.rgb = RGBColor(0,0,0)
        doc.styles['Normal'].font.size = Pt(11)
        doc.styles['Normal'].paragraph_format.line_spacing = 1.2
        paragraphs = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        for sec in doc.sections:
            paragraphs.extend(sec.header.paragraphs + sec.footer.paragraphs)
        for p in paragraphs:
            p.paragraph_format.widow_control = True
            # Prevent the authoring helper's English front matter alignment affecting prose.
            if len(p.text)>110 and p.style.name=='Normal':
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if p.text.startswith('图 ') and p.style.name.startswith('Heading'):
                p.paragraph_format.page_break_before = True
            for r in p.runs:
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Heiti SC' if p.style.name.startswith(('Title','Heading')) else 'Songti SC')
        if '投稿信' in stem:
            for p in doc.paragraphs:
                if p.style.name != 'Title':
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(10)
        doc.save(path)
        print(path)

if __name__=='__main__': main()
